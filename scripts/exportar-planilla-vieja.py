"""Solo para desarrollo: convierte un .xlsx de KASA NORTE a data/planilla-vieja.json
con la misma forma que devuelve la API de Sheets (valores calculados, fechas como
número de serie, filas sin las celdas vacías del final). Así se prueba la
importación en local sin conectarse a Google.

    pip install openpyxl
    python3 scripts/exportar-planilla-vieja.py "KASA NORTE.xlsx"

data/*.json está en .gitignore: el repo es público, los datos no se suben.
"""
import datetime, json, os, sys, warnings
import openpyxl

warnings.filterwarnings('ignore')
TABS = ['COSTOS', 'PRODUCCION', 'INGRESOS Y EGRESOS', 'CLIENTES', 'Base de datos']
EPOCH = datetime.datetime(1899, 12, 30)

def cell(v):
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        return (v - EPOCH).total_seconds() / 86400
    if isinstance(v, datetime.date):
        return (datetime.datetime(v.year, v.month, v.day) - EPOCH).days
    if isinstance(v, datetime.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    return v

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
out = {}
for t in TABS:
    if t not in wb.sheetnames:
        continue
    rows = []
    for r in wb[t].iter_rows(values_only=True):
        vals = [cell(v) for v in r]
        while vals and vals[-1] == '':
            vals.pop()
        rows.append(vals)
    while rows and not rows[-1]:
        rows.pop()
    out[t] = rows
dest = os.path.join(os.path.dirname(__file__), '..', 'data', 'planilla-vieja.json')
json.dump(out, open(dest, 'w'), ensure_ascii=False)
print('ok', {t: len(v) for t, v in out.items()})
